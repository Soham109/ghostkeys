import json, os
from lib import *
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt, Emu

INK4 = RGBColor(0x3A, 0x39, 0x37)
LIT = RGBColor(0xF2, 0xF0, 0xEC)
EDGE = RGBColor(0x55, 0x53, 0x50)
SIG = json.load(open(os.path.join(HERE, "signal.json")))
M = Inches(0.9)
CW = W - 2 * M
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def morph(s):
    """Morph into this slide (fade fallback for apps without Morph)."""
    old = s._element.find(qn("p:transition"))
    if old is not None:
        s._element.remove(old)
    xml = (f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="{P}">'
           f'<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
           f'<p:transition spd="slow" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" p14:dur="1400"><p159:morph option="byObject"/></p:transition>'
           f'</mc:Choice><mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    el = etree.fromstring(xml)
    timing = s._element.find(qn("p:timing"))
    if timing is not None:
        timing.addprevious(el)
    else:
        s._element.append(el)


def box(s, x, y, w, h, line=EDGE, fill=None, radius=None, width=0.75, name=None, shape=MSO_SHAPE.RECTANGLE, dash=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else shape, x, y, w, h)
    if radius is not None:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(width)
        if dash:
            ln = shp.line._get_or_add_ln()
            d = etree.SubElement(ln, qn("a:prstDash"))
            d.set("val", "dash")
    shp.shadow.inherit = False
    if name:
        shp.name = name
    shp.text_frame.text = ""
    return shp


def seg(s, x1, y1, x2, y2, color=EDGE, width=0.75, dash=False, name=None):
    ln = s.shapes.add_connector(1, x1, y1, x2, y2)
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    if dash:
        l = ln.line._get_or_add_ln()
        d = etree.SubElement(l, qn("a:prstDash"))
        d.set("val", "sysDash")
    if name:
        ln.name = name
    return ln


def dot(s, cx, cy, r, fill=LIT, name=None):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r)
    d.fill.solid()
    d.fill.fore_color.rgb = fill
    d.line.fill.background()
    d.shadow.inherit = False
    if name:
        d.name = name
    return d


def poly(s, pts, color=LIT, width=1.0, name=None):
    fb = s.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    fb.add_line_segments(pts[1:], close=False)
    shp = fb.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def small(s, x, y, t, color=INK3, w=Inches(3), align=PP_ALIGN.LEFT, size=10):
    return text(s, x, y, w, Inches(0.3), t.upper(), size=size, font=MONO, color=color, align=align)


# ------------------------------------------------------------------ laptop schematic (top-down)
def laptop(s, ox, oy, lit=False):
    """Draws the top case. Returns dict of zone shapes. Names are stable so Morph can match them."""
    Wc, Hc = Inches(7.0), Inches(4.7)
    shapes = {}
    shapes["case"] = box(s, ox, oy, Wc, Hc, line=EDGE, radius=0.06, name="!!case")
    # top strip
    ts = box(s, ox + Inches(1.25), oy + Inches(0.25), Inches(4.5), Inches(0.28), line=LIT if lit else INK4,
             fill=RGBColor(0x26, 0x25, 0x24) if lit else None, radius=0.5, name="!!zone-top", dash=not lit)
    shapes["top"] = ts
    # keyboard
    kx, ky, kw, kh = ox + Inches(1.25), oy + Inches(0.72), Inches(4.5), Inches(1.95)
    rows, cols = 6, 14
    gap = Inches(0.05)
    cw_ = (kw - gap * (cols - 1)) / cols
    ch_ = (kh - gap * (rows - 1)) / rows
    keys = []
    for r in range(rows):
        for c in range(cols):
            k = box(s, int(kx + c * (cw_ + gap)), int(ky + r * (ch_ + gap)), int(cw_), int(ch_), line=INK4, radius=0.18,
                    name=f"!!key{r}-{c}")
            keys.append(k)
    shapes["keys"] = keys
    # grilles
    for side, gx in (("l", ox + Inches(0.3)), ("r", ox + Wc - Inches(0.3) - Inches(0.7))):
        g = box(s, gx, oy + Inches(0.72), Inches(0.7), Inches(1.95), line=LIT if lit else INK4,
                fill=RGBColor(0x26, 0x25, 0x24) if lit else None, radius=0.2, name=f"!!zone-grille-{side}", dash=not lit)
        shapes["grille-" + side] = g
    # trackpad
    tp = box(s, ox + Inches(2.35), oy + Inches(2.95), Inches(2.3), Inches(1.5), line=INK4, radius=0.08, name="!!trackpad")
    shapes["trackpad"] = tp
    # palm rests
    for side, px in (("l", ox + Inches(0.3)), ("r", ox + Inches(4.85))):
        p = box(s, px, oy + Inches(2.95), Inches(1.85), Inches(1.5), line=LIT if lit else INK4,
                fill=RGBColor(0x26, 0x25, 0x24) if lit else None, radius=0.1, name=f"!!zone-palm-{side}", dash=not lit)
        shapes["palm-" + side] = p
    return shapes


def centre(shp):
    return shp.left + shp.width // 2, shp.top + shp.height // 2


