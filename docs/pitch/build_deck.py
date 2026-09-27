from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
import os

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
SHOTS = "/Users/sohamaggarwal/Desktop/Projects/ghostkeys/app/screenshots"

BG = RGBColor(0x0B, 0x0B, 0x0C)
INK = RGBColor(0xF2, 0xF0, 0xEC)
INK2 = RGBColor(0x9A, 0x98, 0x94)
INK3 = RGBColor(0x5E, 0x5C, 0x59)
LINE = RGBColor(0x2A, 0x29, 0x28)

XL = "Switzer Extralight"
LT = "Switzer Light"
IT = "Switzer Light Italic"
MONO = "Fragment Mono"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    s._anims = []
    return s


def text(s, x, y, w, h, runs, size=20, font=LT, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line=None):
    """runs: str, or list of paragraphs; each paragraph a list of (text, font, color) or a str."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        if spacing is not None and i > 0:
            p.space_before = Pt(spacing)
        parts = para if isinstance(para, list) else [(para, font, color)]
        for t, f, c in parts:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.name = f
            r.font.color.rgb = c
            if f == IT:
                r.font.italic = True
            rPr = r._r.get_or_add_rPr()
            for tag in ("a:latin", "a:ea", "a:cs"):
                pass
    return tb


def label(s, x, y, t, color=INK3, w=Inches(8)):
    return text(s, x, y, w, Inches(0.3), t.upper(), size=10, font=MONO, color=color)


def hline(s, x, y, w):
    ln = s.shapes.add_connector(1, x, y, x + w, y)
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)
    return ln


def vline(s, x, y, h):
    ln = s.shapes.add_connector(1, x, y, x, y + h)
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)
    return ln


def picture(s, path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, x, y, w, h)


def scrim(s, alpha_pct):
    """Flat black veil over a photo (no gradient)."""
    r = s.shapes.add_shape(1, 0, 0, W, H)
    r.line.fill.background()
    r.fill.solid()
    r.fill.fore_color.rgb = BG
    srgb = r.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int(alpha_pct * 1000)))
    return r


# ---------------------------------------------------------------- animation

def anim(s, shape, delay, kind="rise", dur=700):
    s._anims.append((shape.shape_id, delay, kind, dur))


def _effect(ctn_id, spid, delay, kind, dur):
    ids = iter(range(ctn_id, ctn_id + 20))
    preset = {"rise": ("entr", 42), "fade": ("entr", 10), "wipe": ("entr", 22), "settle": ("entr", 10)}[kind][1]
    subtype = 8 if kind == "wipe" else 0
    parts = [f'<p:par><p:cTn id="{next(ids)}" presetID="{preset}" presetClass="entr" presetSubtype="{subtype}" fill="hold" nodeType="withEffect">'
             f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>']
    parts.append(f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                 f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                 f'<p:to><p:strVal val="visible"/></p:to></p:set>')
    if kind == "wipe":
        parts.append(f'<p:animEffect transition="in" filter="wipe(left)"><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur}" decel="100000"/>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')
    else:
        parts.append(f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur}"/>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')
    if kind == "rise":
        parts.append(f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur + 300}" decel="100000" fill="hold"/>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr>'
                     f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="#ppt_y+0.03"/></p:val></p:tav>'
                     f'<p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>')
    if kind == "settle":
        parts.append(f'<p:animScale><p:cBhvr><p:cTn id="{next(ids)}" dur="7000" decel="100000" fill="hold"/>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr>'
                     f'<p:from x="108000" y="108000"/><p:to x="100000" y="100000"/></p:animScale>')
    parts.append('</p:childTnLst></p:cTn></p:par>')
    return "".join(parts), next(ids)


def write_timing(s):
    if not s._anims:
        return
    nid = 5
    effects = []
    for spid, delay, kind, dur in s._anims:
        xml, nid = _effect(nid, spid, delay, kind, dur)
        effects.append(xml)
    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
           f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
           + "".join(effects) +
           f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    s._element.append(etree.fromstring(xml))


def write_transition(s, dur=900):
    xml = f'<p:transition xmlns:p="{P}" spd="slow"><p:fade/></p:transition>'
    t = etree.fromstring(xml)
    # transition must come before timing
    timing = s._element.find(qn("p:timing"))
    if timing is not None:
        timing.addprevious(t)
    else:
        s._element.append(t)


M = Inches(0.9)          # outer margin
CW = W - 2 * M           # content width

# 1 ---------------------------------------------------------------- title
s = new_slide()
bgp = picture(s, os.path.join(IMG, "intro-dark.png"), 0, 0, W, H)
sc = scrim(s, 38)
l1 = label(s, M, Inches(0.8), "Badger BuildFest 2026  ·  Emerging Technology")
t1 = text(s, M, Inches(3.9), CW, Inches(1.6), "Ghostkeys", size=110, font=XL)
t2 = text(s, M, Inches(5.85), CW, Inches(0.7), "Every blank surface is a key.", size=30, font=IT, color=INK)
anim(s, bgp, 0, "settle", 1600)
anim(s, l1, 500, "fade")
anim(s, t1, 700, "rise", 900)
anim(s, t2, 1200, "rise", 900)

# 2 ---------------------------------------------------------------- problem
s = new_slide()
l = label(s, M, Inches(0.8), "01  ·  The problem")
a = text(s, M, Inches(2.2), Inches(10.5), Inches(2.4),
         [[("A MacBook is mostly ", XL, INK), ("dead space.", IT, INK)]], size=72, line=1.0)
b = text(s, M, Inches(4.9), Inches(8), Inches(1.2),
         ["Palm rests. Speaker grilles. The strip above the keys.",
          "None of it does anything, while we spend the day hunting for shortcuts."],
         size=20, color=INK2, spacing=6)
anim(s, l, 200, "fade")
anim(s, a, 400, "rise", 900)
anim(s, b, 1100, "rise")

# 3 ---------------------------------------------------------------- product
s = new_slide()
pic = picture(s, os.path.join(IMG, "zones-dark.png"), Inches(5.4), 0, None, H)
if pic.left + pic.width < W:
    pic.left = W - pic.width
l = label(s, M, Inches(0.8), "02  ·  The product")
a = text(s, M, Inches(1.9), Inches(5), Inches(1.4), [[("Tap it. ", XL, INK), ("It runs.", IT, INK)]], size=54)
rows = [("Left palm rest", "Mute the call"), ("Right grille, twice", "Screenshot"), ("Top strip", "Back to the last app")]
y = Inches(3.7)
items = []
for k, v in rows:
    ln = hline(s, M, y, Inches(4.2))
    tk = text(s, M, y + Inches(0.18), Inches(4.2), Inches(0.35), k.upper(), size=10, font=MONO, color=INK3)
    tv = text(s, M, y + Inches(0.45), Inches(4.2), Inches(0.5), v, size=22)
    items.append((ln, tk, tv))
    y += Inches(1.0)
anim(s, pic, 0, "settle", 1400)
anim(s, l, 200, "fade")
anim(s, a, 400, "rise", 900)
for i, (ln, tk, tv) in enumerate(items):
    anim(s, ln, 1000 + i * 250, "wipe", 600)
    anim(s, tk, 1100 + i * 250, "fade")
    anim(s, tv, 1150 + i * 250, "rise")

# 4 ---------------------------------------------------------------- how
s = new_slide()
l = label(s, M, Inches(0.8), "03  ·  How it works")
a = text(s, M, Inches(1.6), Inches(11), Inches(1.4),
         [[("The sensor is already inside. ", XL, INK), ("We learned to read it.", IT, INK)]], size=44)
cols = [("800", "readings a second", "From the motion sensor in every Apple silicon MacBook. No admin rights, no new hardware."),
        ("33", "measurements per tap", "Two models vote on which spot you touched, and refuse to guess when a tap looks unfamiliar."),
        ("0", "bytes leave the Mac", "Calibration, detection and learning all run on the laptop.")]
cw = CW / 3
anims = []
for i, (n, u, d) in enumerate(cols):
    x = M + int(cw * i)
    if i:
        anims.append(vline(s, x, Inches(3.6), Inches(3.0)))
    pad = Inches(0.35) if i else 0
    tn = text(s, x + pad, Inches(3.5), int(cw) - pad - Inches(0.3), Inches(1.3), n, size=88, font=XL)
    tu = text(s, x + pad, Inches(4.95), int(cw) - pad - Inches(0.3), Inches(0.35), u.upper(), size=10, font=MONO, color=INK3)
    td = text(s, x + pad, Inches(5.4), int(cw) - pad - Inches(0.4), Inches(1.2), d, size=15, color=INK2, line=1.2)
    anims.append((tn, tu, td))
anim(s, l, 200, "fade")
anim(s, a, 400, "rise", 900)
k = 0
for it in anims:
    if isinstance(it, tuple):
        tn, tu, td = it
        anim(s, tn, 1000 + k * 300, "rise", 800)
        anim(s, tu, 1150 + k * 300, "fade")
        anim(s, td, 1250 + k * 300, "fade")
        k += 1
    else:
        anim(s, it, 900 + k * 300, "fade", 500)

# 5 ---------------------------------------------------------------- proof
s = new_slide()
l = label(s, M, Inches(0.8), "04  ·  Measured, not promised")
rows = [("16.5%", "1.4%", "Wrong or false taps"),
        ("2 / min", "0", "False taps while picking up the laptop"),
        ("0%", "84%", "Taps recognised on a lap")]
y = Inches(1.9)
seq = []
for before, after, what in rows:
    ln = hline(s, M, y, CW)
    tw = text(s, M, y + Inches(0.55), Inches(4.3), Inches(0.5), what, size=20, color=INK2)
    tb = text(s, Inches(5.9), y + Inches(0.2), Inches(2.6), Inches(1.1), before, size=54, font=XL, color=INK3)
    ta = text(s, Inches(8.6), y + Inches(0.2), Inches(3.8), Inches(1.1), [[("→  ", XL, INK3), (after, XL, INK)]], size=54)
    seq.append((ln, tw, tb, ta))
    y += Inches(1.45)
foot = text(s, M, Inches(6.55), CW, Inches(0.4),
            "On our own recordings. Every number re-run by an independent reviewer who did not write the code.", size=13, color=INK3)
anim(s, l, 200, "fade")
for i, (ln, tw, tb, ta) in enumerate(seq):
    base = 500 + i * 600
    anim(s, ln, base, "wipe", 700)
    anim(s, tw, base + 150, "fade")
    anim(s, tb, base + 250, "fade")
    anim(s, ta, base + 550, "rise", 800)
anim(s, foot, 2600, "fade")

# 6 ---------------------------------------------------------------- sonar
s = new_slide()
pic = picture(s, os.path.join(IMG, "air-dark.png"), 0, 0, W, H)
sc = scrim(s, 45)
l = label(s, M, Inches(0.8), "05  ·  Beyond taps")
a = text(s, M, Inches(4.2), Inches(11.5), Inches(1.0), [[("Hands in the air. ", XL, INK), ("No camera.", IT, INK)]], size=56)
b = text(s, M, Inches(5.5), Inches(8), Inches(1.0),
         "The speakers play two tones too high to hear. The mic listens to how your hand bends them.",
         size=18, color=INK2, line=1.2)
anim(s, pic, 0, "settle", 1400)
anim(s, l, 300, "fade")
anim(s, a, 600, "rise", 900)
anim(s, b, 1200, "fade")

# 7 ---------------------------------------------------------------- app
s = new_slide()
l = label(s, M, Inches(0.8), "06  ·  The app")
a = text(s, M, Inches(1.6), Inches(4.4), Inches(2.4),
         [[("Calibrates in a minute.", XL, INK)], [("Explains every miss.", IT, INK)]], size=30, line=1.1)
b = text(s, M, Inches(4.0), Inches(4.4), Inches(2.4),
         ["A live view of every tap and where it was felt.",
          "A tester that shows what fired and why.",
          "Per-app layouts for Excel, Figma, Zoom."], size=16, color=INK2, spacing=10)
shot = picture(s, os.path.join(SHOTS, "live.png"), Inches(5.9), Inches(1.35), Inches(6.9))
anim(s, l, 200, "fade")
anim(s, a, 400, "rise", 900)
anim(s, shot, 700, "rise", 1100)
anim(s, b, 1300, "fade")

# 8 ---------------------------------------------------------------- market
s = new_slide()
l = label(s, M, Inches(0.8), "07  ·  Who pays")
a = text(s, M, Inches(1.6), Inches(11.5), Inches(1.4),
         [[("People who live on their laptop ", XL, INK), ("and hate reaching for shortcuts.", IT, INK)]], size=40)
segs = [("Analysts", "Excel all day. Macros on a palm rest."),
        ("Designers", "Tools and undo without leaving the canvas."),
        ("Developers", "Build, run and switch without a chord."),
        ("Anyone on calls", "Mute on a tap they can find blind.")]
cw = CW / 4
seq = []
for i, (t, d) in enumerate(segs):
    x = M + int(cw * i)
    ln = hline(s, x, Inches(3.6), int(cw) - Inches(0.3))
    tt = text(s, x, Inches(3.8), int(cw) - Inches(0.3), Inches(0.5), t, size=22)
    td = text(s, x, Inches(4.35), int(cw) - Inches(0.4), Inches(1.0), d, size=15, color=INK2, line=1.2)
    seq.append((ln, tt, td))
sig = text(s, M, Inches(5.9), CW, Inches(0.9),
           [[("Signal  ", MONO, INK3), ("Mac users already pay for input tools like BetterTouchTool and Raycast. Every Apple silicon MacBook already has the sensor.", LT, INK2)]],
           size=15)
anim(s, l, 200, "fade")
anim(s, a, 400, "rise", 900)
for i, (ln, tt, td) in enumerate(seq):
    anim(s, ln, 1000 + i * 200, "wipe", 600)
    anim(s, tt, 1100 + i * 200, "rise")
    anim(s, td, 1200 + i * 200, "fade")
anim(s, sig, 2100, "fade")

# 9 ---------------------------------------------------------------- business
s = new_slide()
l = label(s, M, Inches(0.8), "08  ·  Business model")
a = text(s, M, Inches(1.6), Inches(11), Inches(1.2), [[("Free to try. ", XL, INK), ("Pay once to keep going.", IT, INK)]], size=44)
tiers = [("Free", "$0", "Forever", "Two hidden buttons on your palm rests."),
         ("Pro", "$29", "Once  ·  $19 at launch", "Every surface, every gesture, every action."),
         ("Teams", "$49", "Per seat, per year", "Pro for the whole desk, managed by IT. From 3 seats.")]
cw = CW / 3
seq = []
for i, (n, p, b_, d) in enumerate(tiers):
    x = M + int(cw * i)
    ln = hline(s, x, Inches(3.3), int(cw) - Inches(0.4))
    tn = text(s, x, Inches(3.5), int(cw), Inches(0.35), n.upper(), size=10, font=MONO, color=INK3)
    tp = text(s, x, Inches(3.85), int(cw), Inches(1.2), p, size=72, font=XL)
    tb = text(s, x, Inches(5.1), int(cw), Inches(0.4), b_, size=14, color=INK2)
    td = text(s, x, Inches(5.55), int(cw) - Inches(0.5), Inches(0.9), d, size=15, color=INK2, line=1.2)
    seq.append((ln, tn, tp, tb, td))
anim(s, l, 200, "fade")
anim(s, a, 400, "rise", 900)
for i, (ln, tn, tp, tb, td) in enumerate(seq):
    base = 1000 + i * 300
    anim(s, ln, base, "wipe", 600)
    anim(s, tn, base + 100, "fade")
    anim(s, tp, base + 150, "rise", 800)
    anim(s, tb, base + 350, "fade")
    anim(s, td, base + 400, "fade")

# 10 --------------------------------------------------------------- why now / edge
s = new_slide()
l = label(s, M, Inches(0.8), "09  ·  Why this wins")
a = text(s, M, Inches(1.6), Inches(11.5), Inches(1.4),
         [[("No hardware. No camera. ", XL, INK), ("No cloud.", IT, INK)]], size=52)
pts = [("The Touch Bar is gone.", "Apple removed the one extra input strip MacBooks had. The space is still there."),
       ("Macro pads cost money and desk space.", "Ghostkeys uses the laptop you already carry, on a train or a couch."),
       ("Trackpad gestures are full.", "Every swipe is taken. The rest of the chassis is not.")]
y = Inches(3.3)
seq = []
for t, d in pts:
    ln = hline(s, M, y, CW)
    tt = text(s, M, y + Inches(0.25), Inches(5.2), Inches(0.6), t, size=21)
    td = text(s, Inches(6.6), y + Inches(0.27), Inches(5.8), Inches(0.8), d, size=16, color=INK2, line=1.2)
    seq.append((ln, tt, td))
    y += Inches(1.1)
anim(s, l, 200, "fade")
anim(s, a, 400, "rise", 900)
for i, (ln, tt, td) in enumerate(seq):
    anim(s, ln, 1000 + i * 350, "wipe", 700)
    anim(s, tt, 1100 + i * 350, "rise")
    anim(s, td, 1250 + i * 350, "fade")

# 11 --------------------------------------------------------------- next
s = new_slide()
l = label(s, M, Inches(0.8), "10  ·  Next 3 to 6 months")
steps = [("Month 1", "Signed public beta from the waitlist."),
         ("Month 2", "Sonar at everyday volume. Swipes in the air."),
         ("Month 3", "Learning from use, proven not to learn the wrong tap."),
         ("Months 4 to 6", "Windows release. Team layouts. First paying teams.")]
y = Inches(1.9)
seq = []
for k_, v in steps:
    ln = hline(s, M, y, CW)
    tk = text(s, M, y + Inches(0.32), Inches(3), Inches(0.4), k_.upper(), size=11, font=MONO, color=INK3)
    tv = text(s, Inches(4.2), y + Inches(0.2), Inches(8.3), Inches(0.7), v, size=26, font=XL)
    seq.append((ln, tk, tv))
    y += Inches(1.15)
anim(s, l, 200, "fade")
for i, (ln, tk, tv) in enumerate(seq):
    anim(s, ln, 500 + i * 350, "wipe", 700)
    anim(s, tk, 600 + i * 350, "fade")
    anim(s, tv, 650 + i * 350, "rise")

# 12 --------------------------------------------------------------- close
s = new_slide()
pic = picture(s, os.path.join(IMG, "try-dark.png"), 0, 0, W, H)
sc = scrim(s, 55)
t1 = text(s, M, Inches(2.3), CW, Inches(1.6), "Ghostkeys", size=96, font=XL, align=PP_ALIGN.CENTER)
t2 = text(s, M, Inches(4.25), CW, Inches(0.7), "Try it on the laptop in front of you.", size=26, font=IT, align=PP_ALIGN.CENTER)
t3 = text(s, M, Inches(5.5), CW, Inches(0.4), "GHOSTKEYS-NINE.VERCEL.APP      GITHUB.COM/SOHAM109/GHOSTKEYS",
          size=11, font=MONO, color=INK2, align=PP_ALIGN.CENTER)
anim(s, pic, 0, "settle", 1600)
anim(s, t1, 500, "rise", 1000)
anim(s, t2, 1100, "rise", 900)
anim(s, t3, 1700, "fade")

for s in prs.slides:
    write_timing(s)
    write_transition(s)

out = os.path.join(HERE, "Ghostkeys-pitch.pptx")
prs.save(out)
print(out)
