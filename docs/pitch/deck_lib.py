from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
import os

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
SHOTS = "/Users/sohamaggarwal/Desktop/Projects/ghostkeys/app/screenshots"

BG = RGBColor(0x0B, 0x0B, 0x0C)
INK = RGBColor(0xF2, 0xF0, 0xEC)
INK2 = RGBColor(0x9A, 0x98, 0x94)
INK3 = RGBColor(0x5E, 0x5C, 0x59)
LINE = RGBColor(0x2A, 0x29, 0x28)

XL = "Switzer Extralight"
LT = "Switzer Light"
IT = "Switzer Light Italic"
MONO = "Fragment Mono"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    s._anims = []
    return s


def text(s, x, y, w, h, runs, size=20, font=LT, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line=None):
    """runs: str, or list of paragraphs; each paragraph a list of (text, font, color) or a str."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        if spacing is not None and i > 0:
            p.space_before = Pt(spacing)
        parts = para if isinstance(para, list) else [(para, font, color)]
        for t, f, c in parts:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.name = f
            r.font.color.rgb = c
            if f == IT:
                r.font.italic = True
            rPr = r._r.get_or_add_rPr()
            for tag in ("a:latin", "a:ea", "a:cs"):
                pass
    return tb


def label(s, x, y, t, color=INK3, w=Inches(8)):
    return text(s, x, y, w, Inches(0.3), t.upper(), size=10, font=MONO, color=color)


def hline(s, x, y, w):
    ln = s.shapes.add_connector(1, x, y, x + w, y)
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)
    return ln


def vline(s, x, y, h):
    ln = s.shapes.add_connector(1, x, y, x, y + h)
    ln.line.color.rgb = LINE
    ln.line.width = Pt(0.75)
    return ln


def picture(s, path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, x, y, w, h)


def scrim(s, alpha_pct):
    """Flat black veil over a photo (no gradient)."""
    r = s.shapes.add_shape(1, 0, 0, W, H)
    r.line.fill.background()
    r.fill.solid()
    r.fill.fore_color.rgb = BG
    srgb = r.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int(alpha_pct * 1000)))
    return r


# ---------------------------------------------------------------- animation

def anim(s, shape, delay, kind="rise", dur=700):
    s._anims.append((shape.shape_id, delay, kind, dur))


def _effect(ctn_id, spid, delay, kind, dur):
    ids = iter(range(ctn_id, ctn_id + 20))
    preset = {"rise": ("entr", 42), "fade": ("entr", 10), "wipe": ("entr", 22), "settle": ("entr", 10)}[kind][1]
    subtype = 8 if kind == "wipe" else 0
    parts = [f'<p:par><p:cTn id="{next(ids)}" presetID="{preset}" presetClass="entr" presetSubtype="{subtype}" fill="hold" nodeType="withEffect">'
             f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>']
    parts.append(f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                 f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                 f'<p:to><p:strVal val="visible"/></p:to></p:set>')
    if kind == "wipe":
        parts.append(f'<p:animEffect transition="in" filter="wipe(left)"><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur}" decel="100000"/>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')
    else:
        parts.append(f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur}"/>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')
    if kind == "rise":
        parts.append(f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur + 300}" decel="100000" fill="hold"/>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr>'
                     f'<p:tavLst><p:tav tm="0"><p:val><p:strVal val="#ppt_y+0.03"/></p:val></p:tav>'
                     f'<p:tav tm="100000"><p:val><p:strVal val="#ppt_y"/></p:val></p:tav></p:tavLst></p:anim>')
    if kind == "settle":
        parts.append(f'<p:animScale><p:cBhvr><p:cTn id="{next(ids)}" dur="7000" decel="100000" fill="hold"/>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr>'
                     f'<p:from x="108000" y="108000"/><p:to x="100000" y="100000"/></p:animScale>')
    parts.append('</p:childTnLst></p:cTn></p:par>')
    return "".join(parts), next(ids)


def write_timing(s):
    if not s._anims:
        return
    nid = 5
    effects = []
    for spid, delay, kind, dur in s._anims:
        xml, nid = _effect(nid, spid, delay, kind, dur)
        effects.append(xml)
    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
           f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
           + "".join(effects) +
           f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    s._element.append(etree.fromstring(xml))


def write_transition(s, dur=900):
    xml = f'<p:transition xmlns:p="{P}" spd="slow"><p:fade/></p:transition>'
    t = etree.fromstring(xml)
    # transition must come before timing
    timing = s._element.find(qn("p:timing"))
    if timing is not None:
        timing.addprevious(t)
    else:
        s._element.append(t)


